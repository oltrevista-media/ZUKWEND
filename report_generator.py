"""
ZUKWEND — Multilingual Billing Agent
Modulo: report_generator.py
Ruolo: Generazione report AR in formato PDF, DOCX, JSON, HTML
Versione: 2.1.5-ENTERPRISE (Full 6-Language Native Edition)
REPTANA MODE | OLTREVISTA 2030 | Offline-First
"""

import html as _html
import io
import json
import logging
import os
import re
from datetime import datetime
from pathlib import Path
from typing import Optional, Union, List
from xml.sax.saxutils import escape as _xml_escape
from i18n_catalog import get_rationale, scan_and_sanitize_text, sanitize_data_recursively, get_legal_disclaimer, get_aging_bucket, format_human_date

try:
    from normalizer import normalize_currency_iso
except ImportError:
    def normalize_currency_iso(c: str) -> str:
        s = str(c or "").strip()
        m = {"€": "EUR", "$": "USD", "£": "GBP", "¥": "JPY", "R$": "BRL"}
        return m.get(s, s.upper() or "EUR")

try:
    from resource_utils import APP_VERSION, get_output_dir, _is_inside_meipass
except ImportError:  # pragma: no cover - esecuzione isolata del modulo
    APP_VERSION = "2.1.5"

    def get_output_dir() -> str:
        folder = os.path.abspath("output")
        os.makedirs(folder, exist_ok=True)
        return folder

    def _is_inside_meipass(path: str) -> bool:
        return False

logger = logging.getLogger("zukwend.report_generator")

_TS_SUFFIX_RE = re.compile(r"_\d{8}_\d{6}(?:_\d+)?$")

# ---------------------------------------------------------------------------
# Sanificazione Percorsi (Clean Path)
# ---------------------------------------------------------------------------
def clean_path(percorsi: Union[str, List, tuple, dict, Path, None], basename_only: bool = False) -> str:
    """
    Sanifica la variabile percorso:
    - Se la variabile e' una lista o tupla, estrae la stringa del primo percorso.
    - Se e' un dict, estrae il primo valore stringa.
    - Rimuove eventuali stringhe grezze tipo "['C:\\...']".
    - Applica os.path.normpath() per rimuovere il raddoppio dei backslash.
    - Se basename_only=True, restituisce direttamente os.path.basename(norm).
    """
    if percorsi is None:
        return ""
    if isinstance(percorsi, Path):
        raw = str(percorsi)
    elif isinstance(percorsi, (list, tuple)):
        raw = str(percorsi[0]) if percorsi else ""
    elif isinstance(percorsi, dict):
        raw = str(next(iter(percorsi.values()))) if percorsi else ""
    else:
        raw = str(percorsi)

    raw = raw.strip()
    if raw.startswith("[") and raw.endswith("]"):
        inner = raw.strip("[]'\" ")
        parts = [p.strip("'\" ") for p in inner.split(",") if p.strip("'\" ")]
        raw = parts[0] if parts else ""

    if not raw:
        return ""

    norm = os.path.normpath(raw)
    return os.path.basename(norm) if basename_only else norm


def _is_lock_error(exc: BaseException) -> bool:
    """
    True per qualsiasi errore di I/O che indica file bloccato / non scrivibile
    (PermissionError, WinError 5 accesso negato, WinError 32 sharing violation, altri OSError).
    """
    return isinstance(exc, (PermissionError, OSError))

def _get_unique_fallback_path(target_path: Path) -> Path:
    """
    Genera un nome file univoco con timestamp in caso di file bloccato o PermissionError/WinError 5.
    Formato: report_AR_INV123_YYYYMMDD_HHMMSS.pdf
    Se il nome contiene gia' un timestamp finale questo viene sostituito (non duplicato).
    """
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    stem = _TS_SUFFIX_RE.sub("", target_path.stem)
    ext = target_path.suffix
    parent = target_path.parent
    parent.mkdir(parents=True, exist_ok=True)

    alt_path = parent / f"{stem}_{ts}{ext}"
    counter = 1
    while alt_path.exists() or alt_path == target_path:
        alt_path = parent / f"{stem}_{ts}_{counter}{ext}"
        counter += 1
    return Path(clean_path(alt_path))

def _write_bytes_with_fallback(path: Path, data: bytes) -> Path:
    """Scrive `data` su `path`; se il file e' bloccato salva con timestamp univoco."""
    path.parent.mkdir(parents=True, exist_ok=True)
    try:
        path.write_bytes(data)
        return Path(clean_path(path))
    except OSError as e:
        alt = _get_unique_fallback_path(path)
        logger.warning(
            f"[LOCK_RECOVERY] Destinazione '{path.name}' bloccata o aperta da altra applicazione ({e}). "
            f"Salvataggio automatico con timestamp univoco: '{alt.name}'"
        )
        alt.write_bytes(data)
        return Path(clean_path(alt))

# ---------------------------------------------------------------------------
# Import librerie opzionali con fallback
# ---------------------------------------------------------------------------
try:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
    from reportlab.platypus.flowables import KeepTogether
    REPORTLAB_AVAILABLE = True
except ImportError:
    REPORTLAB_AVAILABLE = False
    logger.warning("reportlab not available — PDF fallback to plain text")

# ---------------------------------------------------------------------------
# Font Unicode per PDF (accenti à è é ì ò ù ä ö ü ß ñ ç e valute € $ £)
# ---------------------------------------------------------------------------
PDF_FONT_REGULAR = "Helvetica"
PDF_FONT_BOLD = "Helvetica-Bold"
PDF_UNICODE_FONT = False

def _register_pdf_fonts() -> None:
    """Registra una famiglia TrueType Unicode di sistema; fallback su Helvetica (WinAnsi/Latin-1)."""
    global PDF_FONT_REGULAR, PDF_FONT_BOLD, PDF_UNICODE_FONT
    if not REPORTLAB_AVAILABLE:
        return
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    win_fonts = os.path.join(os.environ.get("WINDIR", r"C:\Windows"), "Fonts")
    candidates = [
        (os.path.join(win_fonts, "arial.ttf"), os.path.join(win_fonts, "arialbd.ttf"),
         os.path.join(win_fonts, "ariali.ttf"), os.path.join(win_fonts, "arialbi.ttf")),
        (os.path.join(win_fonts, "segoeui.ttf"), os.path.join(win_fonts, "segoeuib.ttf"),
         os.path.join(win_fonts, "segoeuii.ttf"), os.path.join(win_fonts, "segoeuiz.ttf")),
        ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
         "/usr/share/fonts/truetype/dejavu/DejaVuSans-Oblique.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-BoldOblique.ttf"),
    ]
    for regular, bold, italic, bold_italic in candidates:
        if not (os.path.isfile(regular) and os.path.isfile(bold)):
            continue
        try:
            pdfmetrics.registerFont(TTFont("ZukSans", regular))
            pdfmetrics.registerFont(TTFont("ZukSans-Bold", bold))
            pdfmetrics.registerFont(TTFont("ZukSans-Italic", italic if os.path.isfile(italic) else regular))
            pdfmetrics.registerFont(TTFont("ZukSans-BoldItalic", bold_italic if os.path.isfile(bold_italic) else bold))
            pdfmetrics.registerFontFamily(
                "ZukSans", normal="ZukSans", bold="ZukSans-Bold",
                italic="ZukSans-Italic", boldItalic="ZukSans-BoldItalic",
            )
            PDF_FONT_REGULAR, PDF_FONT_BOLD, PDF_UNICODE_FONT = "ZukSans", "ZukSans-Bold", True
            logger.info(f"Unicode PDF font registered: {os.path.basename(regular)}")
            return
        except Exception as e:
            logger.warning(f"Font registration '{regular}' failed: {e}")
    logger.warning("No TrueType Unicode font available: using Helvetica (WinAnsi Latin-1)")

_register_pdf_fonts()

def _pdf_text_raw(value) -> str:
    """
    Normalizza un testo per il rendering PDF (senza escape XML, per canvas.drawString):
    - rimuove i glifi emoji/simboli non presenti nei font (⚠ U+26A0 + VS16)
    - con il fallback Helvetica converte garantendo preservazione dei caratteri latini/accentati
    """
    s = "" if value is None else str(value)
    s = s.replace("\u26a0\ufe0f", "(!)").replace("\u26a0", "(!)")
    if not PDF_UNICODE_FONT:
        s = s.encode("latin-1", errors="replace").decode("latin-1")
    return s

def _pdf_text(value) -> str:
    """Testo sanitizzato + escape XML (&, <, >) per Paragraph: evita crash del parser su nomi come 'A & B <srl>'."""
    return _xml_escape(_pdf_text_raw(value))

_CTRL_CHARS_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]")

def _clean_strings(obj):
    """Rimuove ricorsivamente i caratteri di controllo non ammessi in XML (DOCX) / PDF."""
    if isinstance(obj, str):
        return _CTRL_CHARS_RE.sub("", obj)
    if isinstance(obj, dict):
        return {k: _clean_strings(v) for k, v in obj.items()}
    if isinstance(obj, list):
        return [_clean_strings(v) for v in obj]
    return obj


try:
    from docx import Document as DocxDocument
    from docx.shared import Pt, RGBColor, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False
    logger.warning("python-docx not available — DOCX disabled")

# ---------------------------------------------------------------------------
# Costanti e Dizionari Multilingua Completi (6 Lingue: IT, EN, DE, FR, ES, PT)
# ---------------------------------------------------------------------------
BRAND_COLOR_HEX = "#FF6B00"   # Arancione OLTREVISTA
TEXT_COLOR      = "#1A1A1A"
LIGHT_GRAY      = "#F5F5F5"
MEDIUM_GRAY     = "#CCCCCC"

DISCLAIMER_TEXT = {
    "it": "Il software fornisce suggerimenti, non decisioni. Output da verificare prima dell'uso operativo. Nessun dato inviato a server esterni.",
    "en": "This software provides suggestions, not decisions. Output must be verified before operational use. No data sent to external servers.",
    "de": "Diese Software liefert Vorschläge, keine Entscheidungen. Die Ausgabe muss vor dem operativen Einsatz geprüft werden. Keine Daten an externe Server gesendet.",
    "fr": "Ce logiciel fournit des suggestions, pas des décisions. La sortie doit être vérifiée avant utilisation opérationnelle. Aucune donnée envoyée à des serveurs externes.",
    "es": "Este software proporciona sugerencias, no decisiones. La salida debe verificarse antes del uso operativo. Ningún dato enviado a servidores externos.",
    "pt": "Este software fornece sugestões, não decisões. A saída deve ser verificada antes do uso operacional. Nenhum dado enviado a servidores externos.",
}

STATUS_LABELS = {
    "it": {"OVERDUE": "SCADUTA", "PENDING": "IN ATTESA", "PAID": "PAGATA", "DISPUTED": "CONTESTATA", "SETTLED": "SALDATA", "CREDIT_NOTE": "NOTA DI CREDITO", "UNKNOWN": "N/D"},
    "en": {"OVERDUE": "OVERDUE", "PENDING": "PENDING", "PAID": "PAID", "DISPUTED": "DISPUTED", "SETTLED": "SETTLED", "CREDIT_NOTE": "CREDIT NOTE", "UNKNOWN": "N/A"},
    "de": {"OVERDUE": "ÜBERFÄLLIG", "PENDING": "AUSSTEHEND", "PAID": "BEZAHLT", "DISPUTED": "BEANSTANDET", "SETTLED": "BEGLICHEN", "CREDIT_NOTE": "GUTSCHRIFT", "UNKNOWN": "N/B"},
    "fr": {"OVERDUE": "EN RETARD", "PENDING": "EN ATTENTE", "PAID": "PAYÉE", "DISPUTED": "CONTESTÉE", "SETTLED": "RÉGLÉE", "CREDIT_NOTE": "AVOIR", "UNKNOWN": "N/D"},
    "es": {"OVERDUE": "VENCIDA", "PENDING": "PENDIENTE", "PAID": "PAGADA", "DISPUTED": "EN DISPUTA", "SETTLED": "Factura Pagada / Settled", "CREDIT_NOTE": "Nota de Crédito / Credit Note", "UNKNOWN": "N/D"},
    "pt": {"OVERDUE": "VENCIDA", "PENDING": "PENDENTE", "PAID": "PAGA", "DISPUTED": "CONTESTADA", "SETTLED": "LIQUIDADA", "CREDIT_NOTE": "NOTA DE CRÉDITO", "UNKNOWN": "N/D"},
}

RISK_LABELS = {
    "it": {"LOW": "BASSO", "MEDIUM": "MEDIO", "HIGH": "ALTO", "NONE": "NESSUNO", "UNKNOWN": "N/D"},
    "en": {"LOW": "LOW", "MEDIUM": "MEDIUM", "HIGH": "HIGH", "NONE": "NONE", "UNKNOWN": "N/A"},
    "de": {"LOW": "NIEDRIG", "MEDIUM": "MITTEL", "HIGH": "HOCH", "NONE": "KEIN", "UNKNOWN": "N/B"},
    "fr": {"LOW": "FAIBLE", "MEDIUM": "MOYEN", "HIGH": "ÉLEVÉ", "NONE": "AUCUN", "UNKNOWN": "N/D"},
    "es": {"LOW": "BAJO", "MEDIUM": "MEDIO", "HIGH": "ALTO", "NONE": "NINGUNO", "UNKNOWN": "N/D"},
    "pt": {"LOW": "BAIXO", "MEDIUM": "MÉDIO", "HIGH": "ALTO", "NONE": "NENHUM", "UNKNOWN": "N/D"},
}

REPORT_LABELS = {
    "it": {
        "report_title": "Rapporto Analisi Aging Crediti",
        "title": "Dati Fattura",
        "inv_num": "Numero Fattura",
        "issue_date": "Data Emissione",
        "due_date": "Data Scadenza",
        "amount": "Importo",
        "client": "Cliente",
        "client_address": "Indirizzo",
        "status": "Stato Fattura",
        "overdue": "Giorni Scaduta",
        "risk": "Rischio Inesigibilità",
        "reminder_lvl": "Livello Sollecito",
        "rem_title": "Sollecito Generato",
        "recs_title": "Azioni Raccomandate",
        "audit_title": "Riepilogo Elaborazione",
        "no_anomalies": "Nessuna anomalia rilevata.",
        "processed": "Data Generazione",
        "offline_tag": "Dati inviati esternamente: NO | Offline-First",
        "field": "Campo",
        "value": "Valore",
        "notice_prefix": "Livello",
        "software_name": "ZUKWEND — Agente di Fatturazione Multilingua",
        "demo_banner": "",
    },
    "en": {
        "report_title": "Accounts Receivable Aging Report",
        "title": "Invoice Details",
        "inv_num": "Invoice Number",
        "issue_date": "Issue Date",
        "due_date": "Due Date",
        "amount": "Amount",
        "client": "Client",
        "client_address": "Address",
        "status": "Invoice Status",
        "overdue": "Days Overdue",
        "risk": "Bad Debt Risk",
        "reminder_lvl": "Reminder Level",
        "rem_title": "Generated Collection Notice",
        "recs_title": "Recommended Actions",
        "audit_title": "Processing Audit",
        "no_anomalies": "No anomalies detected.",
        "processed": "Generation Date",
        "offline_tag": "Data sent externally: NO | Offline-First",
        "field": "Field",
        "value": "Value",
        "notice_prefix": "Level",
        "software_name": "ZUKWEND — Multilingual Billing Agent",
        "demo_banner": "",
    },
    "de": {
        "report_title": "Debitorenalter-Analysebericht",
        "title": "Rechnungsdaten",
        "inv_num": "Rechnungsnummer",
        "issue_date": "Rechnungsdatum",
        "due_date": "Fälligkeitsdatum",
        "amount": "Betrag",
        "client": "Kunde",
        "client_address": "Adresse",
        "status": "Rechnungsstatus",
        "overdue": "Tage überfällig",
        "risk": "Ausfallrisiko",
        "reminder_lvl": "Mahnstufe",
        "rem_title": "Erstelltes Mahnschreiben",
        "recs_title": "Empfohlene Massnahmen",
        "audit_title": "Verarbeitungsprotokoll",
        "no_anomalies": "Keine Unregelmäßigkeiten festgestellt.",
        "processed": "Erstellungsdatum",
        "offline_tag": "Daten extern gesendet: NEIN | Offline-First",
        "field": "Feld",
        "value": "Wert",
        "notice_prefix": "Stufe",
        "software_name": "ZUKWEND — Mehrsprachiger Rechnungsagent",
        "demo_banner": "",
    },
    "fr": {
        "report_title": "Rapport d'Analyse de Solvabilité",
        "title": "Détails de la Facture",
        "inv_num": "Numéro de Facture",
        "issue_date": "Date d'Émission",
        "due_date": "Date d'Échéance",
        "amount": "Montant",
        "client": "Client",
        "client_address": "Adresse",
        "status": "Statut de la Facture",
        "overdue": "Jours de Retard",
        "risk": "Risque d'Insolvabilité",
        "reminder_lvl": "Niveau de Relance",
        "rem_title": "Relance Générée",
        "recs_title": "Actions Recommandées",
        "audit_title": "Journal de Traitement",
        "no_anomalies": "Aucune anomalie détectée.",
        "processed": "Date de Génération",
        "offline_tag": "Données envoyées à l'extérieur : NON | Offline-First",
        "field": "Champ",
        "value": "Valeur",
        "notice_prefix": "Niveau",
        "software_name": "ZUKWEND — Agent de Facturation Multilingue",
        "demo_banner": "",
    },
    "es": {
        "report_title": "Informe de Análisis de Vencimientos",
        "title": "Datos de la Factura",
        "inv_num": "Número de Factura",
        "issue_date": "Fecha de Emisión",
        "due_date": "Fecha de Vencimiento",
        "amount": "Importe",
        "client": "Cliente",
        "client_address": "Dirección",
        "status": "Estado de la Factura",
        "overdue": "Días Vencida",
        "risk": "Riesgo de Impago",
        "reminder_lvl": "Nivel de Reclamación",
        "rem_title": "Aviso de Pago Generado",
        "recs_title": "Acciones Recomendadas",
        "audit_title": "Registro de Procesamiento",
        "no_anomalies": "No se detectaron anomalías.",
        "processed": "Fecha de Generación",
        "offline_tag": "Datos enviados externamente: NO | Offline-First",
        "field": "Campo",
        "value": "Valor",
        "notice_prefix": "Nivel",
        "software_name": "ZUKWEND — Agente de Facturación Multilingüe",
        "demo_banner": "",
    },
    "pt": {
        "report_title": "Relatório de Análise de Aging",
        "title": "Dados da Fatura",
        "inv_num": "Número da Fatura",
        "issue_date": "Data de Emissão",
        "due_date": "Data de Vencimento",
        "amount": "Valor",
        "client": "Cliente",
        "client_address": "Endereço",
        "status": "Estado da Fatura",
        "overdue": "Dias em Atraso",
        "risk": "Risco de Inadimplência",
        "reminder_lvl": "Nível de Cobrança",
        "rem_title": "Aviso de Cobrança Gerado",
        "recs_title": "Ações Recomendadas",
        "audit_title": "Registo de Processamento",
        "no_anomalies": "Nenhuma anomalia detetada.",
        "processed": "Data de Geração",
        "offline_tag": "Dados enviados externamente: NÃO | Offline-First",
        "field": "Campo",
        "value": "Valor",
        "notice_prefix": "Nível",
        "software_name": "ZUKWEND — Agente de Faturação Multilíngue",
        "demo_banner": "",
    },
}

RECOMMENDATIONS_I18N = {
    "it": {
        "threshold": "Importo superiore alla soglia di conformità: richiesta autorizzazione.",
        "legal": "Sollecito livello 3: consultare l'ufficio legale prima di un'azione formale.",
        "overdue_90": "Scaduta da oltre 90 giorni: valutare recupero crediti o svalutazione.",
        "kyc": "Importo elevato: verificare le procedure antiriciclaggio KYC/AML.",
        "none": "Nessuna anomalia rilevata."
    },
    "en": {
        "threshold": "Amount exceeds compliance threshold: authorization required.",
        "legal": "Notice level 3: consult legal counsel prior to formal escalation.",
        "overdue_90": "Overdue exceeds 90 days: evaluate debt recovery or bad debt provisioning.",
        "kyc": "High value amount: verify KYC/AML procedures.",
        "none": "No anomalies detected."
    },
    "de": {
        "threshold": "Betrag überschreitet Compliance-Schwelle: Genehmigung erforderlich.",
        "legal": "Mahnstufe 3: Vor formellen Schritten Rechtsberatung hinzuziehen.",
        "overdue_90": "Über 90 Tage überfällig: Inkassoverfahren oder Wertberichtigung prüfen.",
        "kyc": "Hoher Betrag: KYC/AML-Compliance-Verfahren überprüfen.",
        "none": "Keine Unregelmäßigkeiten festgestellt."
    },
    "fr": {
        "threshold": "Montant supérieur au seuil de conformité : autorisation requise.",
        "legal": "Relance niveau 3 : consulter le service juridique avant toute procédure formelle.",
        "overdue_90": "Retard supérieur à 90 jours : évaluer le recouvrement ou la provision pour créance douteuse.",
        "kyc": "Montant élevé : vérifier les procédures de conformité KYC/AML.",
        "none": "Aucune anomalie détectée."
    },
    "es": {
        "threshold": "Importe superior al umbral de cumplimiento: autorización requerida.",
        "legal": "Nivel de reclamación 3: consultar con asesoría jurídica antes de proceder formalmente.",
        "overdue_90": "Vencimiento superior a 90 días: evaluar cobro de deudas o provisión por insolvencia.",
        "kyc": "Importe elevado: verificar procedimientos KYC/AML.",
        "none": "No se detectaron anomalías."
    },
    "pt": {
        "threshold": "Valor excede o limite de conformidade: autorização necessária.",
        "legal": "Aviso nível 3: consultar o departamento jurídico antes de qualquer ação formal.",
        "overdue_90": "Vencimento superior a 90 dias: avaliar recuperação de crédito ou imparidade.",
        "kyc": "Valor elevado: verificar procedimentos KYC/AML.",
        "none": "Nenhuma anomalia detetada."
    }
}

# ---------------------------------------------------------------------------
# Funzione principale: generate()
# ---------------------------------------------------------------------------
def generate(
    invoice: dict,
    classification: dict,
    compliance: dict,
    billing: dict,
    reminder_text: str,
    language: str = "it",
    output_dir: Optional[str] = None,
    formats: Optional[list] = None,
    is_demo: bool = False,
) -> dict:
    if formats is None:
        formats = ["pdf", "xlsx", "json"]

    lang_code = language.lower() if language else "en"
    if lang_code not in REPORT_LABELS:
        lang_code = "en"

    default_output = get_output_dir()
    if output_dir and not _is_inside_meipass(str(output_dir)):
        out_dir = Path(clean_path(output_dir))
    else:
        out_dir = Path(clean_path(default_output))
    out_dir.mkdir(parents=True, exist_ok=True)

    inv_num_raw = str(invoice.get("invoice_number") or "UNKNOWN")
    inv_num = re.sub(r"[^\w.\-]+", "-", inv_num_raw).strip("-.") or "UNKNOWN"
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    base_name = f"report_AR_{inv_num}_{timestamp}"

    paths = {}
    errors = []

    report_data = _clean_strings(
        _build_report_data(invoice, classification, compliance, billing, reminder_text, lang_code, is_demo)
    )

    for fmt in formats:
        fmt = fmt.lower().strip()
        try:
            if fmt == "pdf":
                p = _generate_pdf(report_data, out_dir / f"{base_name}.pdf", is_demo)
                paths["pdf"] = clean_path(p)
            elif fmt == "xlsx":
                p = _generate_xlsx(report_data, out_dir / f"{base_name}.xlsx")
                paths["xlsx"] = clean_path(p)
            elif fmt == "docx":
                p = _generate_docx(report_data, out_dir / f"{base_name}.docx", is_demo)
                paths["docx"] = clean_path(p)
            elif fmt == "json":
                p = _generate_json(report_data, out_dir / f"{base_name}.json")
                paths["json"] = clean_path(p)
            elif fmt == "html":
                p = _generate_html(report_data, out_dir / f"{base_name}.html", is_demo)
                paths["html"] = clean_path(p)
            else:
                logger.warning(f"Unsupported format: {fmt}")
                errors.append(f"unsupported_format:{fmt}")
        except (PermissionError, OSError) as e:
            if _is_lock_error(e):
                logger.error(
                    f"[LOCK_ERROR] Permission error or file locked (WinError 5 / PermissionError) "
                    f"during generation of {fmt}: {e}."
                )
                errors.append(f"file_locked:{fmt}:{e}")
            else:
                logger.error(f"Error generating {fmt}: {e}")
                errors.append(f"{fmt}:{e}")
        except Exception as e:
            logger.error(f"Error generating {fmt}: {e}")
            errors.append(f"{fmt}:{e}")

    clean_out_dir = clean_path(out_dir)
    logger.info(f"Reports generated: {list(paths.keys())} in {clean_out_dir}")
    return {
        "paths": paths,
        "base_name": base_name,
        "output_dir": clean_out_dir,
        "errors": errors
    }

def _clean_layout_tokens(raw: Optional[str]) -> str:
    if not raw:
        return ""
    s = str(raw).strip()
    pattern = r"(?i)\s*[\-_|/:]*\s*\b(?:invoice|facture|rechnung|page|fattura|fatura|description)\b.*$"
    s = re.sub(pattern, "", s)
    s = re.sub(r"(?i)^\s*description[\s:]*", "", s)
    s = re.sub(r"^[\s:,\-;|]+|[\s:,\-;|]+$", "", s)
    return re.sub(r"\s+", " ", s).strip()

def sanitize_reminder_text(text: str, deadline_days: int = 7) -> str:
    if not text:
        return ""
    clean = text.replace("**", "")
    from datetime import timedelta
    deadline_date = (datetime.now() + timedelta(days=deadline_days)).strftime("%Y-%m-%d")
    clean = re.sub(r"\[insert\s+deadline[^\]]*\]", deadline_date, clean, flags=re.IGNORECASE)
    clean = re.sub(r"\[insert\s+due\s+date[^\]]*\]", deadline_date, clean, flags=re.IGNORECASE)
    clean = re.sub(r"\[insert\s+date[^\]]*\]", deadline_date, clean, flags=re.IGNORECASE)
    clean = re.sub(r"\[insert\s+[^\]]+\]", "", clean, flags=re.IGNORECASE)
    clean = re.sub(r"\[inserisci\s+[^\]]+\]", "", clean, flags=re.IGNORECASE)

    # Rimuovi qualsiasi nota legale / disclaimer dal testo del sollecito
    clean = re.sub(r"(?i)\n*\s*(?:NOTA|NOTE|HINWEIS|AVERTISSEMENT|AVISO)\s*:\s*(?:Questo documento|This document|Dieses Dokument|Ce document|Este documento).*$", "", clean, flags=re.DOTALL)
    # Rimuovi firme fittizie residue in calce (es. Cordiali saluti, The Management, ecc.)
    clean = re.sub(r"(?i)\n*\s*(?:Cordiali saluti|Distinti saluti|Kind regards|Best regards|Mit freundlichen Grüßen|Cordialement|Atentamente|Com os melhores cumprimentos)\s*,?\s*\n+.*$", "", clean, flags=re.DOTALL)
    clean = re.sub(r"(?i)\n*\s*(?:The Management|La Direzione|Die Geschäftsleitung|La Direction|La Dirección|A Direção)\s*$", "", clean)

    clean = re.sub(r"[ \t]+", " ", clean)
    clean = re.sub(r"\n\s*\n\s*\n+", "\n\n", clean)
    return clean.strip()

# ---------------------------------------------------------------------------
# Builder dati report
# ---------------------------------------------------------------------------
def _build_report_data(invoice, classification, compliance, billing, reminder_text, language, is_demo):
    lang_code = (language or "en").lower()
    if lang_code not in REPORT_LABELS:
        lang_code = "en"

    labels = REPORT_LABELS.get(lang_code, REPORT_LABELS["en"])

    risk_map = RISK_LABELS.get(lang_code, RISK_LABELS["en"])
    risk_raw = str(classification.get("risk_level", "UNKNOWN")).upper()
    risk_label = risk_map.get(risk_raw, risk_raw)

    status_map = STATUS_LABELS.get(lang_code, STATUS_LABELS["en"])
    status_raw = str(classification.get("payment_status", "UNKNOWN")).upper()
    status_label = status_map.get(status_raw, status_raw)

    safe_amount = float(invoice.get("amount") or 0.0)

    recommendations = compliance.get("recommendations", [])
    rec_dict = RECOMMENDATIONS_I18N.get(lang_code, RECOMMENDATIONS_I18N["en"])
    translated_recs = []
    for rec in recommendations:
        rec_str = str(rec).strip()
        if "superiore alla soglia" in rec_str or "exceeds compliance threshold" in rec_str or "Compliance-Schwelle" in rec_str:
            translated_recs.append(rec_dict["threshold"])
        elif ("consultare" in rec_str or "consult legal" in rec_str or "Rechtsberatung" in rec_str) and ("legale" in rec_str or "diffida" in rec_str or "3" in rec_str):
            translated_recs.append(rec_dict["legal"])
        elif "90" in rec_str and ("giorni" in rec_str or "days" in rec_str or "Tage" in rec_str or "jours" in rec_str or "días" in rec_str or "dias" in rec_str):
            translated_recs.append(rec_dict["overdue_90"])
        elif "KYC" in rec_str or "AML" in rec_str or "Importo elevato" in rec_str or "High value" in rec_str:
            translated_recs.append(rec_dict["kyc"])
        elif "Nessuna anomalia" in rec_str or "No anomalies" in rec_str or "Keine Unregelmäßigkeiten" in rec_str:
            translated_recs.append(rec_dict["none"])
        else:
            translated_recs.append(rec_str)

    if not translated_recs:
        translated_recs = [rec_dict["none"]]

    try:
        from custom_branding import get_branding
        branding = get_branding()
    except Exception:
        branding = {}

    client_name = _clean_layout_tokens(invoice.get("client_name") or "N/A")
    client_address = _clean_layout_tokens(invoice.get("client_address") or invoice.get("address") or "")
    vendor_name = branding.get("company_name") or _clean_layout_tokens(invoice.get("vendor_name") or "N/A")

    clean_ts = datetime.now().strftime("%Y-%m-%d %H:%M")
    clean_reminder = sanitize_reminder_text(str(reminder_text or ""))

    rationale_val = str(classification.get("classification_rationale") or "")
    if classification.get("rationale_key"):
        rationale_val = get_rationale(
            classification.get("rationale_key"),
            lang_code,
            **(classification.get("rationale_params") or {})
        )
    elif lang_code != "it":
        rationale_val, _ = scan_and_sanitize_text(rationale_val, lang_code)

    canonical_flags = [
        "critical_overdue_bad_debt_risk" if f == "critical_overdue_risk_inesigibilita" else f
        for f in compliance.get("flags", [])
    ]

    # Zero level sync rigoroso
    rem_lvl_raw = classification.get("reminder_level")
    if (
        rem_lvl_raw is None
        or classification.get("action") == "NO_REMINDER_NEEDED"
        or billing.get("action") == "NO_REMINDER_NEEDED"
        or status_raw in ("SETTLED", "CREDIT_NOTE", "PAID", "DISPUTED")
    ):
        final_rem_level = 0
    else:
        try:
            final_rem_level = int(rem_lvl_raw)
        except (ValueError, TypeError):
            final_rem_level = 0

    report_dict = {
        "generated_at": clean_ts,
        "software": labels["software_name"],
        "version": APP_VERSION,
        "language": lang_code,
        "labels": labels,
        "is_demo": is_demo,
        "branding": branding,
        "invoice": {
            "number":         str(invoice.get("invoice_number") or "N/A"),
            "issue_date":     format_human_date(invoice.get("issue_date"), lang_code),
            "due_date":       format_human_date(invoice.get("due_date"), lang_code),
            "issue_date_iso": str(invoice.get("issue_date") or "N/A"),
            "due_date_iso":   str(invoice.get("due_date") or "N/A"),
            "amount":         safe_amount,
            "currency":       normalize_currency_iso(invoice.get("currency") or "EUR"),
            "client":         client_name or "N/A",
            "client_address": client_address,
            "client_vat":     str(invoice.get("client_vat") or invoice.get("vat_number") or ""),
            "vendor":         vendor_name or "N/A",
        },
        "classification": {
            "payment_status":     status_label,
            "payment_status_raw": status_raw,
            "days_overdue":       int(classification.get("days_overdue") or 0),
            "risk_level_raw":     risk_raw,
            "risk_level_label":   risk_label,
            "reminder_level":     final_rem_level,
            "compliance_flag":    bool(classification.get("compliance_flag") or False),
            "rationale":          rationale_val,
            "next_escalation":    str(classification.get("next_escalation_date") or "N/A"),
        },
        "compliance": {
            "flags":           canonical_flags,
            "recommendations": translated_recs,
            "disclaimer":       str(compliance.get("disclaimer") or get_legal_disclaimer(lang_code)),
        },
        "billing": {
            "action":     str(billing.get("action") or "N/A"),
            "channel":    str(billing.get("recommended_channel") or "N/A"),
            "grace_days": int(billing.get("grace_days") or 5),
        },
        "reminder_text": clean_reminder,
        "audit": {
            "processed_at": clean_ts,
            "software_version": APP_VERSION,
            "is_demo": is_demo,
            "data_sent_externally": False,
            "disclaimer": get_legal_disclaimer(lang_code),
        },
    }

    if lang_code != "it":
        report_dict = sanitize_data_recursively(report_dict, lang_code)

    return report_dict

# ---------------------------------------------------------------------------
# Generatore PDF (ReportLab con fallback testuale UTF-8)
# ---------------------------------------------------------------------------
def _generate_pdf(data: dict, path: Path, is_demo: bool = False) -> Path:
    if REPORTLAB_AVAILABLE:
        return _generate_pdf_reportlab(data, path, is_demo)
    else:
        return _generate_pdf_text_fallback(data, path, is_demo)

def _generate_pdf_reportlab(data: dict, path: Path, is_demo: bool) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)

    styles = getSampleStyleSheet()
    for _st in styles.byName.values():
        _is_bold = "bold" in str(getattr(_st, "fontName", "")).lower()
        _st.fontName = PDF_FONT_BOLD if _is_bold else PDF_FONT_REGULAR

    brand_color = colors.HexColor(BRAND_COLOR_HEX)
    labels = data["labels"]
    story = []

    title_style = ParagraphStyle(
        "ZTitle", parent=styles["Title"],
        textColor=brand_color, fontSize=16, spaceAfter=4,
        fontName=PDF_FONT_BOLD
    )

    branding = data.get("branding", {})
    # Logo aziendale se presente
    if branding.get("logo_path") and os.path.exists(branding["logo_path"]):
        try:
            from reportlab.platypus import Image as RLImage
            story.append(RLImage(branding["logo_path"], width=4.5*cm, height=1.8*cm, kind="proportional"))
            story.append(Spacer(1, 0.2*cm))
        except Exception as e_img:
            logger.warning(f"Impossibile caricare logo aziendale: {e_img}")

    story.append(Paragraph(_pdf_text(labels["software_name"]), title_style))
    story.append(Paragraph(_pdf_text(f"{labels['report_title']} | {data['generated_at'][:10]}"), styles["Normal"]))
    
    # Intestazione mittente personalizzata (Custom Branding)
    if branding.get("company_name"):
        comp_info = [branding.get("company_name", "")]
        if branding.get("address"): comp_info.append(branding["address"])
        if branding.get("iban"): comp_info.append(f"IBAN: {branding['iban']}")
        if branding.get("email"): comp_info.append(branding["email"])
        story.append(Paragraph(_pdf_text(" | ".join(comp_info)), styles["Normal"]))

    story.append(HRFlowable(width="100%", thickness=1, color=brand_color, spaceAfter=8))

    inv = data["invoice"]
    clf = data["classification"]
    if labels["title"].strip().upper() != "DESCRIPTION":
        story.append(Paragraph(_pdf_text(labels["title"]), styles["Heading2"]))
    table_data = [
        [labels["inv_num"],      inv["number"]],
        [labels["issue_date"],   inv["issue_date"]],
        [labels["due_date"],     inv["due_date"]],
        [labels["amount"],       f"{inv['currency']} {inv['amount']:,.2f}"],
        [labels["client"],       inv["client"]],
    ]
    if inv.get("client_address"):
        table_data.append([labels["client_address"], inv["client_address"]])
    table_data.extend([
        [labels["status"],       clf["payment_status"]],
        [labels["overdue"],      str(clf["days_overdue"])],
        [labels["risk"],         clf["risk_level_label"]],
        [labels["reminder_lvl"], f"{labels['notice_prefix']} {clf['reminder_level']}"],
    ])
    cell_style = ParagraphStyle("ZCell", parent=styles["Normal"], fontSize=10, leading=12)
    table_cells = [[Paragraph(_pdf_text(k), cell_style), Paragraph(_pdf_text(v), cell_style)]
                   for k, v in table_data]
    t = Table(table_cells, colWidths=[5.5*cm, 10.5*cm])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (0,-1), colors.HexColor(LIGHT_GRAY)),
        ("FONTNAME", (0,0), (-1,-1), PDF_FONT_REGULAR),
        ("FONTSIZE", (0,0), (-1,-1), 10),
        ("GRID", (0,0), (-1,-1), 0.5, colors.HexColor(MEDIUM_GRAY)),
        ("VALIGN", (0,0), (-1,-1), "TOP"),
        ("ROWBACKGROUNDS", (0,0), (-1,-1), [colors.white, colors.HexColor(LIGHT_GRAY)]),
    ]))
    story.append(t)
    story.append(Spacer(1, 0.5*cm))

    story.append(Paragraph(_pdf_text(labels["rem_title"]), styles["Heading2"]))
    for line in data["reminder_text"].split("\n"):
        line_clean = line.replace("**", "").strip()
        if line_clean and line_clean.upper() != "DESCRIPTION":
            story.append(Paragraph(_pdf_text(line_clean), styles["Normal"]))
        elif not line_clean:
            story.append(Spacer(1, 0.2*cm))
    story.append(Spacer(1, 0.5*cm))

    if data["compliance"]["recommendations"]:
        story.append(Paragraph(_pdf_text(labels["recs_title"]), styles["Heading2"]))
        for rec in data["compliance"]["recommendations"]:
            rec_clean = str(rec).replace("**", "").strip()
            story.append(Paragraph(_pdf_text(f"• {rec_clean}"), styles["Normal"]))
        story.append(Spacer(1, 0.3*cm))

    story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor(MEDIUM_GRAY), spaceAfter=4))
    audit_style = ParagraphStyle("Audit", parent=styles["Normal"], fontSize=8, textColor=colors.gray)
    clean_proc_at = data["audit"]["processed_at"][:16].replace("T", " ")
    story.append(Paragraph(_pdf_text(f"{labels['processed']}: {clean_proc_at} | ZUKWEND — Multilingual Billing Agent 2.1.5"), audit_style))
    story.append(Paragraph(_pdf_text(labels["offline_tag"]), audit_style))
    story.append(Spacer(1, 0.2*cm))

    disc_style = ParagraphStyle("Disc", parent=styles["Normal"], fontSize=8,
                                 textColor=colors.gray, borderPadding=4)
    story.append(Paragraph(_pdf_text(data["audit"]["disclaimer"]), disc_style))

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=A4,
        leftMargin=1.5*cm, rightMargin=1.5*cm,
        topMargin=1.5*cm, bottomMargin=1.5*cm
    )
    doc.build(story)
    pdf_bytes = buffer.getvalue()
    saved_path = _write_bytes_with_fallback(path, pdf_bytes)
    logger.info(f"PDF generated: {clean_path(saved_path)}")
    return Path(clean_path(saved_path))

def _generate_pdf_text_fallback(data: dict, path: Path, is_demo: bool) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    txt_path = path.with_suffix(".txt")
    content = _build_text_report(data, is_demo)
    target_path = txt_path
    try:
        target_path.write_text(content, encoding="utf-8")
    except (PermissionError, OSError) as e:
        if _is_lock_error(e):
            target_path = _get_unique_fallback_path(txt_path)
            logger.warning(
                f"[LOCK_RECOVERY] TXT destination '{txt_path.name}' locked ({e}). "
                f"Saving to '{target_path.name}'"
            )
            target_path.write_text(content, encoding="utf-8")
        else:
            raise e
    logger.warning(f"PDF fallback (text): {clean_path(target_path)}")
    return Path(clean_path(target_path))

def _build_text_report(data: dict, is_demo: bool) -> str:
    sep = "=" * 60
    inv = data["invoice"]
    clf = data["classification"]
    labels = data["labels"]
    lines = [
        sep,
        labels["software_name"],
        f"{labels['report_title']} | {data['generated_at'][:10]}",
        sep,
        labels["title"].upper(),
        f"  {labels['inv_num']}:     {inv['number']}",
        f"  {labels['due_date']}:    {inv['due_date']}",
        f"  {labels['amount']}:      {inv['currency']} {inv['amount']:,.2f}",
        f"  {labels['client']}:      {inv['client']}",
    ] + ([f"  {labels['client_address']}:     {inv['client_address']}"] if inv.get("client_address") else []) + [
        f"  {labels['status']}:      {clf['payment_status']} ({clf['days_overdue']})",
        f"  {labels['risk']}:        {clf['risk_level_label']} | {labels['notice_prefix']} {clf['reminder_level']}",
        sep,
        labels["rem_title"].upper(),
        data["reminder_text"],
        sep,
        labels["recs_title"].upper(),
    ] + [f"  • {r}" for r in data["compliance"]["recommendations"]] + [
        sep,
        f"{labels['processed']}: {data['audit']['processed_at']}",
        data["audit"]["disclaimer"],
        sep,
    ]
    return "\n".join(lines)

# ---------------------------------------------------------------------------
# Generatore DOCX
# ---------------------------------------------------------------------------
def _generate_docx(data: dict, path: Path, is_demo: bool = False) -> Path:
    if not DOCX_AVAILABLE:
        raise RuntimeError("python-docx non disponibile")

    labels = data["labels"]
    doc = DocxDocument()
    inv = data["invoice"]
    clf = data["classification"]

    title = doc.add_heading(labels["software_name"], level=1)
    title.runs[0].font.color.rgb = RGBColor(0xFF, 0x6B, 0x00)
    doc.add_paragraph(f"{labels['report_title']} | {data['generated_at'][:10]}")

    doc.add_heading(labels["title"], level=2)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = labels["field"]
    hdr[1].text = labels["value"]
    rows = [
        (labels["inv_num"],      inv["number"]),
        (labels["issue_date"],   str(inv["issue_date"])),
        (labels["due_date"],     str(inv["due_date"])),
        (labels["amount"],       f"{inv['currency']} {inv['amount']:,.2f}"),
        (labels["client"],       inv["client"]),
    ] + ([(labels["client_address"], inv["client_address"])] if inv.get("client_address") else []) + [
        (labels["status"],       clf["payment_status"]),
        (labels["overdue"],      str(clf["days_overdue"])),
        (labels["risk"],         clf["risk_level_label"]),
        (labels["reminder_lvl"], f"{labels['notice_prefix']} {clf['reminder_level']}"),
    ]
    for field, value in rows:
        row = table.add_row().cells
        row[0].text = field
        row[1].text = str(value)

    doc.add_heading(labels["rem_title"], level=2)
    doc.add_paragraph(data["reminder_text"])

    doc.add_heading(labels["recs_title"], level=2)
    for rec in data["compliance"]["recommendations"]:
        doc.add_paragraph(f"• {rec}")

    doc.add_heading(labels["audit_title"], level=3)
    audit_para = doc.add_paragraph(
        f"{labels['processed']}: {data['audit']['processed_at']} | {data['software']}\n"
        f"{labels['offline_tag']}"
    )
    audit_para.runs[0].font.size = Pt(9)

    disclaimer = doc.add_paragraph(data["audit"]["disclaimer"])
    disclaimer.runs[0].font.size = Pt(9)
    disclaimer.runs[0].font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    path.parent.mkdir(parents=True, exist_ok=True)
    target_path = path
    try:
        doc.save(str(target_path))
    except (PermissionError, OSError) as e:
        if _is_lock_error(e):
            target_path = _get_unique_fallback_path(path)
            logger.warning(
                f"[LOCK_RECOVERY] DOCX destination '{path.name}' locked ({e}). "
                f"Saving to '{target_path.name}'"
            )
            doc.save(str(target_path))
        else:
            raise e
    clean_docx = clean_path(target_path)
    logger.info(f"DOCX generated: {clean_docx}")
    return Path(clean_docx)

# ---------------------------------------------------------------------------
# Generatore JSON
# ---------------------------------------------------------------------------
def _generate_json(data: dict, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    days_od = data["classification"]["days_overdue"]
    bucket_label = get_aging_bucket(days_od)

    json_out = {
        "zukwend_version": data["version"],
        "generated_at": data["generated_at"],
        "language": data["language"],
        "report_title": data["labels"]["report_title"],
        "invoice": data["invoice"],
        "classification": data["classification"],
        "compliance": data["compliance"],
        "billing": data["billing"],
        "reminder_text": data["reminder_text"],
        "audit": data["audit"],
        "erp_ingestion": {
            "erp_compatible_systems": ["SAP", "Zucchetti", "TeamSystem"],
            "invoice_number": data["invoice"]["number"],
            "client": data["invoice"]["client"],
            "client_address": data["invoice"].get("client_address", ""),
            "issue_date": data["invoice"].get("issue_date_iso", data["invoice"]["issue_date"]),
            "due_date": data["invoice"].get("due_date_iso", data["invoice"]["due_date"]),
            "currency": data["invoice"]["currency"],
            "amount": data["invoice"]["amount"],
            "payment_status": data["classification"]["payment_status"],
            "days_overdue": days_od,
            "aging_bucket": bucket_label,
            "risk_level": data["classification"]["risk_level_label"],
            "reminder_level": data["classification"]["reminder_level"],
            "recommended_action": data["billing"]["action"],
            "export_timestamp": data["audit"]["processed_at"]
        }
    }
    if data.get("branding"):
        json_out["branding"] = data["branding"]

    content = json.dumps(json_out, indent=2, ensure_ascii=False)
    if data.get("language") and data["language"] != "it":
        sanitized_content, leaks = scan_and_sanitize_text(content, data["language"])
        if leaks:
            logger.warning(f"[ANTI-LEAK] Sanitized Italian leaks in JSON: {leaks}")
        content = sanitized_content
    target_path = path
    try:
        target_path.write_text(content, encoding="utf-8")
    except (PermissionError, OSError) as e:
        if _is_lock_error(e):
            target_path = _get_unique_fallback_path(path)
            logger.warning(
                f"[LOCK_RECOVERY] JSON destination '{path.name}' locked ({e}). "
                f"Saving to '{target_path.name}'"
            )
            target_path.write_text(content, encoding="utf-8")
        else:
            raise e
    clean_json = clean_path(target_path)
    logger.info(f"JSON generated: {clean_json}")
    return Path(clean_json)

# ---------------------------------------------------------------------------
# Generatore Excel (.xlsx) Formattato con Aging Buckets (0-30/31-60/61-90/90+)
# ---------------------------------------------------------------------------
def _generate_xlsx(data: dict, path: Path) -> Path:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "AR Aging Analysis"
    ws.views.sheetView[0].showGridLines = True

    # Palette colori professionali
    brand_fill = PatternFill(start_color="FF6B00", end_color="FF6B00", fill_type="solid")
    dark_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    gray_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    highlight_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")

    font_title = Font(name="Calibri", size=15, bold=True, color="FFFFFF")
    font_section = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_bold = Font(name="Calibri", size=11, bold=True)
    font_normal = Font(name="Calibri", size=11)
    font_muted = Font(name="Calibri", size=9, italic=True, color="64748B")

    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1")
    )

    inv = data["invoice"]
    clf = data["classification"]
    labels = data["labels"]
    branding = data.get("branding", {})

    # 1. Header Banner
    ws.merge_cells("A1:G1")
    title_cell = ws["A1"]
    title_cell.value = f"ZUKWEND — {labels.get('report_title', 'Accounts Receivable Aging Report')}"
    title_cell.font = font_title
    title_cell.fill = brand_fill
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 36

    # 2. Informazioni Mittente (Custom Branding)
    row = 3
    if branding.get("company_name"):
        cred_lbl = "CREDITORE:" if data.get("language") == "it" else "CREDITOR:"
        ws.cell(row=row, column=1, value=cred_lbl).font = font_bold
        ws.cell(row=row, column=2, value=branding.get("company_name", "")).font = font_bold
        if branding.get("address"):
            row += 1
            ws.cell(row=row, column=1, value="ADDRESS:").font = font_muted
            ws.cell(row=row, column=2, value=branding.get("address", "")).font = font_normal
        if branding.get("iban"):
            row += 1
            ws.cell(row=row, column=1, value="IBAN:").font = font_muted
            ws.cell(row=row, column=2, value=branding.get("iban", "")).font = font_normal
        if branding.get("email"):
            row += 1
            ws.cell(row=row, column=1, value="EMAIL:").font = font_muted
            ws.cell(row=row, column=2, value=branding.get("email", "")).font = font_normal
        row += 1

    # 3. Sezione Dettagli Fattura
    ws.cell(row=row, column=1, value="INVOICE & DEBTOR DETAILS").font = font_section
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
    for c in range(1, 8):
        ws.cell(row=row, column=c).fill = dark_fill
    ws.row_dimensions[row].height = 24
    row += 1

    invoice_fields = [
        (labels.get("inv_num", "Invoice Number"), inv["number"]),
        (labels.get("issue_date", "Issue Date"), inv["issue_date"]),
        (labels.get("due_date", "Due Date"), inv["due_date"]),
        (labels.get("amount", "Amount"), f"{inv['currency']} {inv['amount']:,.2f}"),
        (labels.get("client", "Client"), inv["client"]),
        (labels.get("client_address", "Client Address"), inv.get("client_address", "")),
        (labels.get("status", "Status"), clf["payment_status"]),
        (labels.get("overdue", "Days Overdue"), clf["days_overdue"]),
        (labels.get("risk", "Risk Level"), clf["risk_level_label"]),
        (labels.get("reminder_lvl", "Reminder Level"), f"{labels.get('notice_prefix', 'Level')} {clf['reminder_level']}"),
    ]

    for label, val in invoice_fields:
        c1 = ws.cell(row=row, column=1, value=label)
        c2 = ws.cell(row=row, column=2, value=str(val))
        c1.font = font_bold
        c1.fill = gray_fill
        c2.font = font_normal
        c1.border = thin_border
        c2.border = thin_border
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
        for cc in range(2, 5):
            ws.cell(row=row, column=cc).border = thin_border
        row += 1

    row += 1

    # 4. Aging Buckets Section (0-30, 31-60, 61-90, 90+)
    ws.cell(row=row, column=1, value="AR AGING BUCKETS BREAKDOWN").font = font_section
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
    for c in range(1, 8):
        ws.cell(row=row, column=c).fill = dark_fill
    ws.row_dimensions[row].height = 24
    row += 1

    bucket_headers = ["Current / Not Due", "0-30 Days", "31-60 Days", "61-90 Days", "90+ Days (Critical)", "Total Overdue"]
    for col_idx, b_title in enumerate(bucket_headers, start=1):
        cell = ws.cell(row=row, column=col_idx, value=b_title)
        cell.font = font_header
        cell.fill = brand_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    ws.row_dimensions[row].height = 22
    row += 1

    days_od = clf["days_overdue"]
    amt = inv["amount"]

    b_not_due = amt if days_od <= 0 else 0.0
    b_0_30    = amt if 1 <= days_od <= 30 else 0.0
    b_31_60   = amt if 31 <= days_od <= 60 else 0.0
    b_61_90   = amt if 61 <= days_od <= 90 else 0.0
    b_90_plus = amt if days_od > 90 else 0.0
    total_od  = amt if days_od > 0 else 0.0

    bucket_values = [b_not_due, b_0_30, b_31_60, b_61_90, b_90_plus, total_od]
    for col_idx, b_val in enumerate(bucket_values, start=1):
        cell = ws.cell(row=row, column=col_idx, value=b_val)
        cell.font = font_bold if col_idx in (5, 6) else font_normal
        cell.number_format = f'"{inv["currency"]} " #,##0.00'
        cell.alignment = Alignment(horizontal="right", vertical="center")
        cell.border = thin_border
        if b_val > 0 and col_idx in (5, 6):
            cell.fill = highlight_fill
    row += 2

    # 5. Scheda 2: Ingestione ERP (SAP, Zucchetti, TeamSystem)
    ws_erp = wb.create_sheet(title="ERP_Ingestion")
    ws_erp.views.sheetView[0].showGridLines = True
    erp_cols = [
        "INVOICE_ID", "CLIENT_NAME", "CLIENT_ADDRESS", "ISSUE_DATE", "DUE_DATE",
        "CURRENCY", "AMOUNT", "PAYMENT_STATUS", "DAYS_OVERDUE", "AGING_BUCKET",
        "RISK_LEVEL", "REMINDER_LEVEL", "RECOMMENDED_ACTION", "EXPORT_TIMESTAMP"
    ]
    for c_idx, h_text in enumerate(erp_cols, start=1):
        c = ws_erp.cell(row=1, column=c_idx, value=h_text)
        c.font = font_header
        c.fill = dark_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border

    bucket_str = get_aging_bucket(days_od)

    erp_row_data = [
        inv["number"], inv["client"], inv.get("client_address", ""), inv.get("issue_date_iso", inv["issue_date"]), inv.get("due_date_iso", inv["due_date"]),
        inv["currency"], inv["amount"], clf["payment_status"], clf["days_overdue"], bucket_str,
        clf["risk_level_label"], clf["reminder_level"], data["billing"]["action"], data["audit"]["processed_at"]
    ]
    for c_idx, val in enumerate(erp_row_data, start=1):
        c = ws_erp.cell(row=2, column=c_idx, value=val)
        c.font = font_normal
        c.border = thin_border
        if c_idx == 7:  # Amount
            c.number_format = "#,##0.00"
            c.alignment = Alignment(horizontal="right")

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 15)

    for col in ws_erp.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_erp.column_dimensions[col_letter].width = max(max_len + 4, 15)

    target_path = path
    try:
        wb.save(str(target_path))
    except (PermissionError, OSError) as e:
        if _is_lock_error(e):
            target_path = _get_unique_fallback_path(path)
            logger.warning(f"[LOCK_RECOVERY] XLSX destination '{path.name}' locked ({e}). Saving to '{target_path.name}'")
            wb.save(str(target_path))
        else:
            raise e
    clean_xlsx = clean_path(target_path)
    logger.info(f"XLSX generated: {clean_xlsx}")
    return Path(clean_xlsx)

# ---------------------------------------------------------------------------
# Generatore HTML
# ---------------------------------------------------------------------------
def _generate_html(data: dict, path: Path, is_demo: bool = False) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    def _esc_map(src: dict) -> dict:
        return {k: (_html.escape(v, quote=False) if isinstance(v, str) else v) for k, v in src.items()}

    inv = _esc_map(data["invoice"])
    clf = _esc_map(data["classification"])
    labels = _esc_map(data["labels"])
    data = {**data, "audit": _esc_map(data["audit"])}
    risk_class = {"LOW": "risk-low", "MEDIUM": "risk-medium", "HIGH": "risk-high"}.get(
        data["classification"]["risk_level_raw"], ""
    )
    recs_html = "".join(f"<li>{_html.escape(str(r))}</li>" for r in data["compliance"]["recommendations"])
    reminder_html = _html.escape(data["reminder_text"]).replace("\n", "<br>")

    html = f"""<!DOCTYPE html>
<html lang="{data['language']}">
<head>
  <meta charset="UTF-8">
  <meta name="viewport" content="width=device-width, initial-scale=1.0">
  <title>{labels['report_title']} — ZUKWEND</title>
  <style>
    body {{ font-family: Calibri, 'Segoe UI', sans-serif; color: #1a1a1a; margin: 0; padding: 20px; background: #f9f9f9; }}
    .header {{ background: #ff6b00; color: white; padding: 20px 24px; border-radius: 6px 6px 0 0; }}
    .header h1 {{ margin: 0; font-size: 1.4em; }}
    .header p  {{ margin: 4px 0 0; font-size: 0.9em; opacity: 0.9; }}
    .container {{ background: white; border: 1px solid #e0e0e0; border-radius: 0 0 6px 6px; padding: 24px; }}
    .section-title {{ color: #ff6b00; border-bottom: 2px solid #ff6b00; padding-bottom: 4px; margin-top: 24px; font-size: 1.1em; }}
    table {{ width: 100%; border-collapse: collapse; margin-top: 12px; }}
    th, td {{ padding: 8px 12px; text-align: left; border: 1px solid #e0e0e0; font-size: 0.95em; }}
    th {{ background: #f5f5f5; font-weight: bold; width: 35%; }}
    .risk-low    {{ color: #28a745; font-weight: bold; }}
    .risk-medium {{ color: #ffc107; font-weight: bold; }}
    .risk-high   {{ color: #dc3545; font-weight: bold; }}
    .reminder-box {{ background: #f9f4ef; border-left: 4px solid #ff6b00; padding: 16px; margin-top: 12px; border-radius: 0 4px 4px 0; font-size: 0.95em; line-height: 1.6; }}
    .disclaimer {{ font-size: 0.82em; color: #888; margin-top: 20px; border-top: 1px solid #eee; padding-top: 12px; }}
    .audit {{ font-size: 0.82em; color: #aaa; margin-top: 8px; }}
    ul {{ padding-left: 20px; }}
    li {{ margin-bottom: 4px; }}
  </style>
</head>
<body>
  <div class="header">
    <h1>{labels['software_name']}</h1>
    <p>{labels['report_title']} | {data['generated_at'][:10]} | {data['version']}</p>
  </div>
  <div class="container">
    <div class="section-title">{labels['title']}</div>
    <table>
      <tr><th>{labels['inv_num']}</th><td>{inv['number']}</td></tr>
      <tr><th>{labels['issue_date']}</th><td>{inv['issue_date']}</td></tr>
      <tr><th>{labels['due_date']}</th><td>{inv['due_date']}</td></tr>
      <tr><th>{labels['amount']}</th><td><strong>{inv['currency']} {inv['amount']:,.2f}</strong></td></tr>
      <tr><th>{labels['client']}</th><td>{inv['client']}</td></tr>
      {f"<tr><th>{labels['client_address']}</th><td>{inv['client_address']}</td></tr>" if inv.get("client_address") else ""}
      <tr><th>{labels['status']}</th><td>{clf['payment_status']}</td></tr>
      <tr><th>{labels['overdue']}</th><td>{clf['days_overdue']}</td></tr>
      <tr><th>{labels['risk']}</th><td class="{risk_class}">{clf['risk_level_label']}</td></tr>
      <tr><th>{labels['reminder_lvl']}</th><td>{labels['notice_prefix']} {clf['reminder_level']}</td></tr>
    </table>
    <div class="section-title">{labels['rem_title']}</div>
    <div class="reminder-box">{reminder_html}</div>
    <div class="section-title">{labels['recs_title']}</div>
    <ul>{recs_html}</ul>
    <div class="disclaimer">{data['audit']['disclaimer']}</div>
    <div class="audit">
      {labels['processed']}: {data['audit']['processed_at']} | {labels['offline_tag']}
    </div>
  </div>
</body>
</html>"""

    target_path = path
    try:
        target_path.write_text(html, encoding="utf-8")
    except (PermissionError, OSError) as e:
        if _is_lock_error(e):
            target_path = _get_unique_fallback_path(path)
            logger.warning(
                f"[LOCK_RECOVERY] HTML destination '{path.name}' locked ({e}). "
                f"Saving to '{target_path.name}'"
            )
            target_path.write_text(html, encoding="utf-8")
        else:
            raise e
    clean_html = clean_path(target_path)
    logger.info(f"HTML generated: {clean_html}")
    return Path(clean_html)


# ---------------------------------------------------------------------------
# Consolidated Dossier Exports (v2.1.6 Multi-Currency & Strict P.IVA)
# ---------------------------------------------------------------------------

CONSOLIDATED_EMAIL_I18N = {
    "it": {
        "subject": "Estratto Conto e Riepilogo Fatture Scadute - {client_name}",
        "greeting": "Gentile {client_name},",
        "intro": "Dai nostri controlli contabili aggiornati, risultano ad oggi le seguenti posizioni aperte in sospeso:",
        "totals_header": "Riepilogo Saldi Scaduti per Valuta:",
        "inv_line": "  • Fattura n. {number} | Data: {issue_date} | Scadenza: {due_date} | Scaduta da: {days_overdue} gg | Importo: {amount} {currency}",
        "total_line": "  • Totale {currency}: {total} {currency}",
        "action_prompt": "Vi invitiamo a verificare la corrispondenza con i Vostri archivi e a procedere con il saldo del dovuto. Qualora il pagamento fosse già stato effettuato, Vi preghiamo di trasmetterci la relativa contabile bancaria.",
        "bank_header": "Coordinate bancarie per il bonifico:",
        "iban_label": "IBAN:",
    },
    "en": {
        "subject": "Statement of Account & Outstanding Invoices Summary - {client_name}",
        "greeting": "Dear {client_name},",
        "intro": "According to our accounting records, the following invoice(s) are currently outstanding:",
        "totals_header": "Outstanding Balances by Currency:",
        "inv_line": "  • Invoice #{number} | Date: {issue_date} | Due: {due_date} | Overdue: {days_overdue} days | Amount: {amount} {currency}",
        "total_line": "  • Total {currency}: {total} {currency}",
        "action_prompt": "Please review your records and confirm settlement of the open balance. If payment has already been remitted, please provide the remittance confirmation so we may update our records.",
        "bank_header": "Bank Details for Remittance:",
        "iban_label": "IBAN:",
    },
    "de": {
        "subject": "Kontoauszug und Übersicht offener Rechnungen - {client_name}",
        "greeting": "Sehr geehrte Damen und Herren ({client_name}),",
        "intro": "Gemäß unserer Buchhaltung sind aktuell folgende Rechnungen noch zur Zahlung offen:",
        "totals_header": "Offene Beträge nach Währung:",
        "inv_line": "  • Rechnung Nr. {number} | Datum: {issue_date} | Fällig am: {due_date} | Überfällig: {days_overdue} Tage | Betrag: {amount} {currency}",
        "total_line": "  • Gesamt {currency}: {total} {currency}",
        "action_prompt": "Wir bitten Sie, Ihren Kontenstand zu prüfen und den fälligen Ausgleich zu veranlassen. Sollte die Zahlung bereits erfolgt sein, betrachten Sie dieses Schreiben bitte als gegenstandslos.",
        "bank_header": "Bankverbindung:",
        "iban_label": "IBAN:",
    },
    "fr": {
        "subject": "Relevé de compte et récapitulatif des factures en attente - {client_name}",
        "greeting": "Madame, Monsieur ({client_name}),",
        "intro": "D'après nos écritures comptables, les factures suivantes restent à ce jour en attente de règlement :",
        "totals_header": "Récapitulatif des soldes par devise :",
        "inv_line": "  • Facture n° {number} | Date : {issue_date} | Échéance : {due_date} | Retard : {days_overdue} jours | Montant : {amount} {currency}",
        "total_line": "  • Total {currency} : {total} {currency}",
        "action_prompt": "Nous vous remercions de bien vouloir vérifier vos écritures et procéder au règlement du solde ouvert. Si le paiement a déjà été effectué, merci d'ignorer ce rappel.",
        "bank_header": "Coordonnées bancaires pour le virement :",
        "iban_label": "IBAN :",
    },
    "es": {
        "subject": "Estado de Cuenta y Resumen de Facturas Pendientes - {client_name}",
        "greeting": "Estimado/a {client_name},",
        "intro": "Según nuestros registros contables, a la fecha figuran las siguientes facturas pendientes de pago:",
        "totals_header": "Resumen de saldos pendientes por moneda:",
        "inv_line": "  • Factura n.° {number} | Fecha: {issue_date} | Vencimiento: {due_date} | Vencida: {days_overdue} días | Importe: {amount} {currency}",
        "total_line": "  • Total {currency}: {total} {currency}",
        "action_prompt": "Le solicitamos tenga a bien revisar sus registros y gestionar el saldo pendiente a la brevedad. Si ya ha efectuado el pago, le rogamos desestime este mensaje.",
        "bank_header": "Datos bancarios para transferencia:",
        "iban_label": "IBAN:",
    },
    "pt": {
        "subject": "Extrato de Conta e Resumo de Faturas em Aberto - {client_name}",
        "greeting": "Prezado(a) {client_name},",
        "intro": "De acordo com os nossos registos contabilísticos, encontram-se atualmente em aberto as seguintes faturas:",
        "totals_header": "Resumo dos saldos em aberto por moeda:",
        "inv_line": "  • Fatura n.º {number} | Data: {issue_date} | Vencimento: {due_date} | Em atraso: {days_overdue} dias | Valor: {amount} {currency}",
        "total_line": "  • Total {currency}: {total} {currency}",
        "action_prompt": "Agradecemos que verifique os seus registos e proceda à regularização do valor em dívida. Caso o pagamento já tenha sido emitido, queira desconsiderar este aviso.",
        "bank_header": "Dados bancários para transferência:",
        "iban_label": "IBAN:",
    }
}

STATEMENT_TITLES = {
    "it": "ESTRATTO CONTO CLIENTI CONSOLIDATO",
    "en": "CONSOLIDATED STATEMENT OF ACCOUNT",
    "de": "KONSOLIDIERTER KONTOAUSZUG",
    "fr": "RELEVÉ DE COMPTE CONSOLIDÉ",
    "es": "ESTADO DE CUENTA CONSOLIDADO",
    "pt": "EXTRATO DE CONTA CONSOLIDADO",
}

STATEMENT_SUBTITLES = {
    "it": "Dossier Consolidato: {client_name} | Data: {today_str}",
    "en": "Consolidated Dossier: {client_name} | Date: {today_str}",
    "de": "Konsolidierte Kundenakte: {client_name} | Datum: {today_str}",
    "fr": "Dossier consolidé : {client_name} | Date : {today_str}",
    "es": "Expediente consolidado: {client_name} | Fecha: {today_str}",
    "pt": "Dossier consolidado: {client_name} | Data: {today_str}",
}

DAYS_OVERDUE_SUFFIX = {
    "it": "gg",
    "en": "d",
    "de": "T.",
    "fr": "j",
    "es": "días",
    "pt": "dias",
}

TOTAL_LABEL_I18N = {
    "it": "Totale {curr}:",
    "en": "Total {curr}:",
    "de": "Gesamt {curr}:",
    "fr": "Total {curr} :",
    "es": "Total {curr}:",
    "pt": "Total {curr}:",
}


def _extract_invoice_record(item: dict) -> dict:
    """Normalizza un elemento invoice da report o dict grezzo."""
    if isinstance(item, dict) and "invoice" in item:
        inv = item["invoice"]
        clf = item.get("classification", {})
        bill = item.get("billing", {})
    elif isinstance(item, dict) and "data" in item and isinstance(item["data"], dict) and "invoice" in item["data"]:
        inv = item["data"]["invoice"]
        clf = item["data"].get("classification", {})
        bill = item["data"].get("billing", {})
    elif isinstance(item, dict):
        inv = item
        clf = item.get("classification", {})
        bill = item.get("billing", {})
    else:
        inv = {}
        clf = {}
        bill = {}

    num = str(inv.get("number") or inv.get("invoice_number") or "N/A")
    iss_d = str(inv.get("issue_date") or inv.get("issue_date_iso") or "N/A")
    due_d = str(inv.get("due_date") or inv.get("due_date_iso") or "N/A")
    iss_d_iso = str(inv.get("issue_date_iso") or inv.get("issue_date") or "N/A")
    due_d_iso = str(inv.get("due_date_iso") or inv.get("due_date") or "N/A")
    days_od = int(clf.get("days_overdue") or inv.get("days_overdue") or 0)
    st = str(clf.get("payment_status") or clf.get("payment_status_raw") or "OVERDUE")
    amt = float(inv.get("amount") or 0.0)
    curr = normalize_currency_iso(inv.get("currency") or "EUR")
    client = str(inv.get("client") or inv.get("client_name") or "Client")
    client_addr = str(inv.get("client_address") or inv.get("address") or "")
    client_vat = str(inv.get("client_vat") or inv.get("vat_number") or "")
    risk_lbl = str(clf.get("risk_level_label") or clf.get("risk_level_raw") or "MEDIUM")
    rem_lvl = int(clf.get("reminder_level") or 0)
    action = str(bill.get("action") or "SEND_REMINDER")

    return {
        "number": num,
        "issue_date": iss_d,
        "due_date": due_d,
        "issue_date_iso": iss_d_iso,
        "due_date_iso": due_d_iso,
        "days_overdue": days_od,
        "status": st,
        "amount": amt,
        "currency": curr,
        "client": client,
        "client_address": client_addr,
        "client_vat": client_vat,
        "risk_level_label": risk_lbl,
        "reminder_level": rem_lvl,
        "action": action,
    }


def generate_consolidated_statement_pdf(
    client_data: dict,
    invoices: list,
    output_path: Union[str, Path],
    language: str = "en",
    is_demo: bool = False,
    branding: Optional[dict] = None
) -> Path:
    """
    Genera STATEMENT_OF_ACCOUNT_CONSOLIDATED.pdf per il dossier cliente.
    Tabella analitica dettagliata con SUBTOTALI SEPARATI PER VALUTA.
    """
    out_file = Path(output_path)
    out_file.parent.mkdir(parents=True, exist_ok=True)

    lang_code = (language or "en").lower()[:2]
    labels = REPORT_LABELS.get(lang_code, REPORT_LABELS["en"])
    if branding is None:
        try:
            from custom_branding import get_branding
            branding = get_branding()
        except Exception:
            branding = {}

    client_name = client_data.get("client_name") or "CLIENT"
    client_vat = client_data.get("vat") or "N/A"
    today_str = datetime.now().strftime("%Y-%m-%d")

    records = [_extract_invoice_record(i) for i in invoices]

    # Calcolo subtotali separati per valuta (nessuna somma algebrica incrociata)
    subtotals_by_curr = {}
    for r in records:
        c = r["currency"]
        subtotals_by_curr[c] = round(subtotals_by_curr.get(c, 0.0) + r["amount"], 2)

    # Tentativo ReportLab
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import cm
        from reportlab.lib import colors
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

        doc = SimpleDocTemplate(
            str(out_file),
            pagesize=A4,
            leftMargin=1.5 * cm,
            rightMargin=1.5 * cm,
            topMargin=1.5 * cm,
            bottomMargin=1.5 * cm,
        )

        styles = getSampleStyleSheet()
        for _st in styles.byName.values():
            _is_bold = "bold" in str(getattr(_st, "fontName", "")).lower()
            _st.fontName = PDF_FONT_BOLD if _is_bold else PDF_FONT_REGULAR

        brand_col = colors.HexColor(BRAND_COLOR_HEX)
        dark_col = colors.HexColor("#1E293B")
        subtotal_bg = colors.HexColor("#FEF3C7")

        story = []

        # Logo opzionale
        if branding.get("logo_path") and os.path.exists(branding["logo_path"]):
            try:
                from reportlab.platypus import Image as RLImage
                story.append(RLImage(branding["logo_path"], width=4.5 * cm, height=1.6 * cm, kind="proportional"))
                story.append(Spacer(1, 0.2 * cm))
            except Exception as e_img:
                logger.warning(f"Consolidated PDF logo load error: {e_img}")

        title_style = ParagraphStyle(
            "ConsolidatedTitle",
            parent=styles["Title"],
            textColor=brand_col,
            fontSize=16,
            fontName=PDF_FONT_BOLD,
            spaceAfter=3,
        )
        subtitle_style = ParagraphStyle(
            "ConsolidatedSubtitle",
            parent=styles["Normal"],
            textColor=colors.HexColor("#4B5563"),
            fontSize=10,
            spaceAfter=10,
        )
        body_style = ParagraphStyle(
            "ConsolidatedBody",
            parent=styles["Normal"],
            fontSize=9,
            leading=12,
        )
        bold_style = ParagraphStyle(
            "ConsolidatedBold",
            parent=styles["Normal"],
            fontName=PDF_FONT_BOLD,
            fontSize=9,
            leading=12,
        )

        title_text = STATEMENT_TITLES.get(lang_code, STATEMENT_TITLES["en"])
        subtitle_tmpl = STATEMENT_SUBTITLES.get(lang_code, STATEMENT_SUBTITLES["en"])
        subtitle_text = subtitle_tmpl.format(client_name=client_name, today_str=today_str)
        story.append(Paragraph(_pdf_text(f"ZUKWEND — {title_text}"), title_style))
        story.append(Paragraph(_pdf_text(subtitle_text), subtitle_style))

        # Blocco Intestazione: Mittente vs Destinatario
        sender_lines = [
            f"<b>{branding.get('company_name', 'OLTREVISTA / ZUKWEND')}</b>",
            branding.get("address", ""),
            f"IBAN: {branding.get('iban', 'N/A')}" if branding.get("iban") else "",
            f"Email: {branding.get('email', 'N/A')}" if branding.get("email") else "",
        ]
        sender_html = "<br/>".join([l for l in sender_lines if l])

        client_lines = [
            f"<b>{client_name}</b>",
            f"VAT / Tax ID: <b>{client_vat}</b>",
            f"Total Invoices: {len(records)}",
            f"Dossier Key: {client_data.get('folder_slug', 'N/A')}",
        ]
        client_html = "<br/>".join(client_lines)

        meta_table = Table(
            [
                [
                    Paragraph(_pdf_text_raw(sender_html), body_style),
                    Paragraph(_pdf_text_raw(client_html), body_style),
                ]
            ],
            colWidths=[9.0 * cm, 8.5 * cm],
        )
        meta_table.setStyle(
            TableStyle(
                [
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F8FAFC")),
                    ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
                    ("PADDING", (0, 0), (-1, -1), 6),
                ]
            )
        )
        story.append(meta_table)
        story.append(Spacer(1, 0.4 * cm))

        # Tabella analitica fatture
        th_inv = labels.get("inv_num", "Invoice #")
        th_iss = labels.get("issue_date", "Issue Date")
        th_due = labels.get("due_date", "Due Date")
        th_od = labels.get("overdue", "Days Overdue")
        th_st = labels.get("status", "Status")
        th_amt = labels.get("amount", "Amount")

        table_data = [
            [
                Paragraph(_pdf_text(th_inv), bold_style),
                Paragraph(_pdf_text(th_iss), bold_style),
                Paragraph(_pdf_text(th_due), bold_style),
                Paragraph(_pdf_text(th_od), bold_style),
                Paragraph(_pdf_text(th_st), bold_style),
                Paragraph(_pdf_text(th_amt), bold_style),
            ]
        ]

        od_suffix = DAYS_OVERDUE_SUFFIX.get(lang_code, "d")
        for r in records:
            amt_str = f"{r['amount']:,.2f} {r['currency']}"
            table_data.append(
                [
                    Paragraph(_pdf_text(r["number"]), body_style),
                    Paragraph(_pdf_text(r["issue_date"]), body_style),
                    Paragraph(_pdf_text(r["due_date"]), body_style),
                    Paragraph(_pdf_text(f"{r['days_overdue']} {od_suffix}"), body_style),
                    Paragraph(_pdf_text(r["status"]), body_style),
                    Paragraph(_pdf_text(amt_str), bold_style),
                ]
            )

        # Righe subtotali separati per valuta (MAI sommare valute diverse)
        tot_lbl_tmpl = TOTAL_LABEL_I18N.get(lang_code, TOTAL_LABEL_I18N["en"])
        subtotal_row_start = len(table_data)
        for curr, tot in subtotals_by_curr.items():
            tot_str = f"{tot:,.2f} {curr}"
            table_data.append(
                [
                    Paragraph("", body_style),
                    Paragraph("", body_style),
                    Paragraph("", body_style),
                    Paragraph("", body_style),
                    Paragraph(_pdf_text(tot_lbl_tmpl.format(curr=curr)), bold_style),
                    Paragraph(_pdf_text(tot_str), bold_style),
                ]
            )

        inv_table = Table(
            table_data,
            colWidths=[3.6 * cm, 2.7 * cm, 2.7 * cm, 2.3 * cm, 2.8 * cm, 3.4 * cm],
        )

        t_style = [
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E293B")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("PADDING", (0, 0), (-1, -1), 4),
        ]
        for row_idx in range(subtotal_row_start, len(table_data)):
            t_style.append(("BACKGROUND", (0, row_idx), (-1, row_idx), subtotal_bg))
            t_style.append(("LINEABOVE", (0, row_idx), (-1, row_idx), 1, brand_col))

        inv_table.setStyle(TableStyle(t_style))
        story.append(inv_table)
        story.append(Spacer(1, 0.4 * cm))

        # Disclaimer
        discl_text = get_legal_disclaimer(lang_code)
        story.append(
            Paragraph(
                _pdf_text(f"Note: Multi-currency balances are strictly partitioned. {discl_text}"),
                ParagraphStyle("Discl", parent=styles["Normal"], fontSize=7.5, textColor=colors.HexColor("#64748B")),
            )
        )

        doc.build(story)
        logger.info(f"Consolidated Statement PDF created: {out_file}")
        return out_file

    except Exception as e:
        logger.warning(f"ReportLab consolidated PDF generation failed ({e}), writing text fallback.")
        lines = [
            f"=== ZUKWEND {title_text} ===",
            f"Client: {client_name}",
            f"VAT / Tax ID: {client_vat}",
            f"Generated: {today_str}",
            "",
            "INVOICE BREAKDOWN:",
        ]
        od_suffix = DAYS_OVERDUE_SUFFIX.get(lang_code, "d")
        for r in records:
            lines.append(f"- #{r['number']} | Issue: {r['issue_date']} | Due: {r['due_date']} | Overdue: {r['days_overdue']} {od_suffix} | Status: {r['status']} | Amount: {r['amount']:,.2f} {r['currency']}")
        lines.append("")
        lines.append("SUBTOTALS BY CURRENCY:")
        tot_lbl_tmpl = TOTAL_LABEL_I18N.get(lang_code, TOTAL_LABEL_I18N["en"])
        for curr, tot in subtotals_by_curr.items():
            lines.append(f"  * {tot_lbl_tmpl.format(curr=curr)} {tot:,.2f} {curr}")
        out_file.write_text("\n".join(lines), encoding="utf-8")
        return out_file


def generate_consolidated_reminder_email(
    client_data: dict,
    invoices: list,
    output_path: Union[str, Path],
    language: str = "en",
    branding: Optional[dict] = None
) -> Path:
    """
    Genera REMINDER_EMAIL_CONSOLIDATED.txt contenente OGGETTO e CORPO unificato
    con riepilogo analitico e subtotali separati per valuta.
    """
    out_file = Path(output_path)
    out_file.parent.mkdir(parents=True, exist_ok=True)

    lang_code = (language or "en").lower()[:2]
    i18n = CONSOLIDATED_EMAIL_I18N.get(lang_code, CONSOLIDATED_EMAIL_I18N["en"])

    if branding is None:
        try:
            from custom_branding import get_branding
            branding = get_branding()
        except Exception:
            branding = {}

    client_name = client_data.get("client_name") or "Client"
    records = [_extract_invoice_record(i) for i in invoices]

    # Subtotali valuta
    subtotals_by_curr = {}
    for r in records:
        c = r["currency"]
        subtotals_by_curr[c] = round(subtotals_by_curr.get(c, 0.0) + r["amount"], 2)

    subject = i18n["subject"].format(client_name=client_name)
    greeting = i18n["greeting"].format(client_name=client_name)
    intro = i18n["intro"]

    body_lines = [
        f"SUBJECT: {subject}",
        "",
        "BODY:",
        greeting,
        "",
        intro,
        "",
    ]

    inv_tmpl = i18n.get("inv_line", "  • Invoice #{number} | Date: {issue_date} | Due: {due_date} | Overdue: {days_overdue} days | Amount: {amount} {currency}")
    for r in records:
        line = inv_tmpl.format(
            number=r["number"],
            issue_date=r["issue_date"],
            due_date=r["due_date"],
            days_overdue=r["days_overdue"],
            amount=f"{r['amount']:,.2f}",
            currency=r["currency"],
        )
        body_lines.append(line)

    body_lines.append("")
    body_lines.append(i18n["totals_header"])
    tot_tmpl = i18n.get("total_line", "  • Total {currency}: {total} {currency}")
    for curr, tot in subtotals_by_curr.items():
        body_lines.append(tot_tmpl.format(currency=curr, total=f"{tot:,.2f}"))

    body_lines.append("")
    body_lines.append(i18n["action_prompt"])

    if branding.get("iban"):
        body_lines.append("")
        body_lines.append(i18n["bank_header"])
        body_lines.append(f"  {i18n['iban_label']} {branding['iban']}")

    # Clean text: nessun saluto generico ('Kind regards') o firma fittizia ('Accounts Receivable Department')
    content = "\n".join(body_lines)
    out_file.write_text(content, encoding="utf-8")
    logger.info(f"Consolidated Reminder Email generated: {out_file}")
    return out_file

